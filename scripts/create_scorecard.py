import os
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

PREDICTIONS_FILE = os.environ.get(
    "ADP_PREDICTIONS_FILE"
)

CONFIG_FILE = os.environ.get(
    "ADP_CONFIG_FILE"
)

OUTPUT_FILE = os.environ.get(
    "ADP_SCORECARD_FILE"
)


# ============================================================
# START
# ============================================================

print("=" * 60)
print("CREATE SCORECARD")
print("=" * 60)


# ============================================================
# CHECK FILES
# ============================================================

if not PREDICTIONS_FILE:
    raise ValueError(
        "Prediction file path was not provided."
    )

if not CONFIG_FILE:
    raise ValueError(
        "Config file path was not provided."
    )

if not OUTPUT_FILE:
    raise ValueError(
        "Scorecard output path was not provided."
    )


if not os.path.exists(PREDICTIONS_FILE):

    raise FileNotFoundError(
        f"Predictions file not found: "
        f"{PREDICTIONS_FILE}"
    )


if not os.path.exists(CONFIG_FILE):

    raise FileNotFoundError(
        f"Config file not found: "
        f"{CONFIG_FILE}"
    )


# ============================================================
# READ PREDICTIONS
# ============================================================

predictions = pd.read_csv(
    PREDICTIONS_FILE
)


# ============================================================
# KEEP ONLY IMPORTANT COLUMNS
# ============================================================

prediction_columns = [
    "video",
    "prediction",
    "confidence",
    "accident_probability",
    "timestamp_seconds"
]


missing_columns = [
    column
    for column in prediction_columns
    if column not in predictions.columns
]


if missing_columns:

    raise ValueError(
        "Missing prediction columns: "
        + ", ".join(missing_columns)
    )


prediction_sheet = predictions[
    prediction_columns
].copy()


# ============================================================
# RENAME COLUMNS
# ============================================================

prediction_sheet = prediction_sheet.rename(
    columns={
        "video": "Video",
        "prediction": "Prediction",
        "confidence": "Confidence",
        "accident_probability":
            "Accident Probability",
        "timestamp_seconds":
            "Timestamp (seconds)"
    }
)


# ============================================================
# READ CONFIGURATION
# ============================================================

with open(
    CONFIG_FILE,
    "r",
    encoding="utf-8"
) as file:

    config = json.load(file)


# ============================================================
# CONFIGURATION SHEET
# ============================================================

config_items = [

    (
        "Model Version",
        config.get("model_version")
    ),

    (
        "FPS",
        config.get("fps")
    ),

    (
        "Resolution",
        "x".join(
            map(
                str,
                config.get(
                    "image_size",
                    []
                )
            )
        )
    ),

    (
        "Sequence Length",
        config.get(
            "sequence_length"
        )
    ),

    (
        "Threshold",
        config.get(
            "threshold"
        )
    )

]


config_sheet = pd.DataFrame(
    config_items,
    columns=[
        "Configuration",
        "Value"
    ]
)


# ============================================================
# CREATE OUTPUT DIRECTORY
# ============================================================

output_directory = os.path.dirname(
    OUTPUT_FILE
)

if output_directory:

    os.makedirs(
        output_directory,
        exist_ok=True
    )


# ============================================================
# CREATE EXCEL FILE
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    prediction_sheet.to_excel(
        writer,
        sheet_name="Video Results",
        index=False
    )

    config_sheet.to_excel(
        writer,
        sheet_name="Configuration",
        index=False
    )


# ============================================================
# FORMAT EXCEL
# ============================================================

workbook = load_workbook(
    OUTPUT_FILE
)


for worksheet in workbook.worksheets:

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


    # --------------------------------------------------------
    # COLUMN WIDTH
    # --------------------------------------------------------

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = (
            get_column_letter(
                column_cells[0].column
            )
        )

        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            45
        )


    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"


# ============================================================
# SAVE
# ============================================================

workbook.save(
    OUTPUT_FILE
)


# ============================================================
# COMPLETED
# ============================================================

print()
print("=" * 60)
print("SCORECARD CREATED SUCCESSFULLY")
print("=" * 60)

print(
    f"\nSaved to:\n{OUTPUT_FILE}"
)

print(
    "\nSheets created:"
)

print(
    "1. Video Results"
)

print(
    "2. Configuration"
)